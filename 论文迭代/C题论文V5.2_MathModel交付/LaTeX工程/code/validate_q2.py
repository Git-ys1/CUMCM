from __future__ import annotations
import csv,json
from pathlib import Path
import numpy as np
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"outputs"/"q2"
def validate()->dict:
    with (OUT/"solution.csv").open(encoding="utf-8-sig",newline="") as f:
        rows=list(csv.DictReader(f))
    if len(rows)!=334*144: raise ValueError(f"solution rows: {len(rows)}")
    grid=np.asarray([float(r["grid_plan_kwh"]) for r in rows]).reshape(334,144)
    charge=np.asarray([float(r["charge_actual_kwh"]) for r in rows]).reshape(334,144)
    discharge=np.asarray([float(r["discharge_actual_kwh"]) for r in rows]).reshape(334,144)
    emergency=np.asarray([float(r["emergency_kwh"]) for r in rows]).reshape(334,144)
    with (OUT/"daily_metrics.csv").open(encoding="utf-8-sig",newline="") as f:
        daily=list(csv.DictReader(f))
    wb=openpyxl.load_workbook(OUT/"result2.xlsx",data_only=True,read_only=True)
    template=openpyxl.load_workbook(ROOT.parent/"CUMCM2026Problems"/"C题"/"附件"/"附件5"/"result2.xlsx",
                                   data_only=True,read_only=True)
    plan_rows=list(wb.worksheets[0].iter_rows(min_row=2,max_row=335,min_col=2,max_col=147,values_only=True))
    plan_values=np.asarray([r[:144] for r in plan_rows],dtype=float)
    total_values=np.asarray([r[144] for r in plan_rows],dtype=float)
    cost_values=np.asarray([r[145] for r in plan_rows],dtype=float)
    storage_rows=list(wb.worksheets[1].iter_rows(min_row=2,max_row=2005,min_col=1,max_col=6,values_only=True))
    storage_charge=np.asarray([float(r[2]) for r in storage_rows]).reshape(334,6)
    storage_discharge=np.asarray([float(r[3]) for r in storage_rows]).reshape(334,6)
    storage_initial=np.asarray([float(storage_rows[d*6][5]) for d in range(334)])
    storage_terminal=np.asarray([float(storage_rows[d*6+1][5]) for d in range(334)])
    expected_charge=charge.reshape(334,6,24).sum(axis=2)
    expected_discharge=discharge.reshape(334,6,24).sum(axis=2)
    emergency_rows=list(wb.worksheets[2].iter_rows(min_row=2,min_col=1,max_col=3,values_only=True))
    emergency_sheet_total=float(sum(float(r[2] or 0) for r in emergency_rows))
    expected_initial=np.asarray([float(r["soc_initial"]) for r in daily])
    expected_terminal=np.asarray([float(r["soc_terminal"]) for r in daily])
    expected_cost=np.asarray([float(r["plan_cost"]) for r in daily])
    result={
        "sheet_names_preserved":wb.sheetnames==template.sheetnames,
        "solution_rows":len(rows),"plan_shape":list(plan_values.shape),"storage_rows":len(storage_rows),
        "max_plan_diff":float(np.max(np.abs(plan_values-grid))),
        "max_plan_total_diff":float(np.max(np.abs(total_values-grid.sum(axis=1)))),
        "max_plan_cost_diff":float(np.max(np.abs(cost_values-expected_cost))),
        "max_storage_charge_diff":float(np.max(np.abs(storage_charge-expected_charge))),
        "max_storage_discharge_diff":float(np.max(np.abs(storage_discharge-expected_discharge))),
        "max_storage_initial_soc_diff":float(np.max(np.abs(storage_initial-expected_initial))),
        "max_storage_terminal_soc_diff":float(np.max(np.abs(storage_terminal-expected_terminal))),
        "emergency_sheet_total_kwh":emergency_sheet_total,
        "emergency_csv_total_kwh":float(emergency.sum()),
        "emergency_total_diff":abs(emergency_sheet_total-float(emergency.sum())),
    }
    result["passed"]=bool(result["sheet_names_preserved"] and result["plan_shape"]==[334,144] and
        len(storage_rows)==2004 and max(result[k] for k in result if k.startswith("max_"))<1e-7 and
        result["emergency_total_diff"]<1e-7)
    (OUT/"xlsx_verification.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    metrics=json.loads((OUT/"metrics.json").read_text(encoding="utf-8"))
    metrics["xlsx_passed"]=result["passed"]
    (OUT/"metrics.json").write_text(json.dumps(metrics,ensure_ascii=False,indent=2),encoding="utf-8")
    return result
if __name__=="__main__": print(json.dumps(validate(),ensure_ascii=False,indent=2))
