from __future__ import annotations
import csv,json
from pathlib import Path
import numpy as np
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT.parent/"CUMCM2026Problems"/"C题"/"附件"/"附件5"
def _csv(path:Path)->list[dict]:
    with path.open(encoding="utf-8-sig",newline="") as f: return list(csv.DictReader(f))
def _maxdiff(a,b)->float: return float(np.max(np.abs(np.asarray(a,dtype=float)-np.asarray(b,dtype=float))))
def validate_q1()->dict:
    folder=ROOT/"outputs"/"q1"; rows=_csv(folder/"solution.csv")
    wb=openpyxl.load_workbook(folder/"result1.xlsx",data_only=True,read_only=True)
    template=openpyxl.load_workbook(DATA/"result1.xlsx",data_only=True,read_only=True)
    grid=np.asarray([float(r["grid_kwh"]) for r in rows])
    xgrid=np.asarray(list(wb.worksheets[0].iter_rows(min_row=2,max_row=145,min_col=2,max_col=2,values_only=True)),dtype=float).ravel()
    result={"rows":len(rows),"sheet_names_preserved":wb.sheetnames==template.sheetnames,
            "max_grid_diff":_maxdiff(grid,xgrid)}
    result["passed"]=bool(len(rows)==144 and result["sheet_names_preserved"] and result["max_grid_diff"]<1e-7)
    return result
def validate_annual(folder:Path,xlsx_name:str,kind:str)->dict:
    rows=_csv(folder/"solution.csv"); daily=_csv(folder/"daily_metrics.csv")
    if len(rows)!=334*144 or len(daily)!=334: raise ValueError(f"{folder}: row count mismatch")
    wb=openpyxl.load_workbook(folder/xlsx_name,data_only=True,read_only=True)
    template=openpyxl.load_workbook(DATA/xlsx_name,data_only=True,read_only=True)
    has_adjust=kind=="q3"; storage_index=2 if has_adjust else 1; emergency_index=3 if has_adjust else 2
    plan_field="baseline_grid_kwh" if has_adjust else "grid_plan_kwh"
    plan=np.asarray([float(r[plan_field]) for r in rows]).reshape(334,144)
    charge=np.asarray([float(r["charge_actual_kwh"]) for r in rows]).reshape(334,144)
    discharge=np.asarray([float(r["discharge_actual_kwh"]) for r in rows]).reshape(334,144)
    emergency=np.asarray([float(r["emergency_kwh"]) for r in rows]).reshape(334,144)
    plan_rows=list(wb.worksheets[0].iter_rows(min_row=2,max_row=335,min_col=2,max_col=147,values_only=True))
    xplan=np.asarray([r[:144] for r in plan_rows],dtype=float)
    xtotal=np.asarray([r[144] for r in plan_rows],dtype=float)
    xcost=np.asarray([r[145] for r in plan_rows],dtype=float)
    expected_plan_cost=np.asarray([float(r["baseline_cost"] if has_adjust else r["plan_cost"]) for r in daily])
    result={"solution_rows":len(rows),"daily_rows":len(daily),"sheet_names_preserved":wb.sheetnames==template.sheetnames,
            "max_plan_diff":_maxdiff(xplan,plan),"max_plan_total_diff":_maxdiff(xtotal,plan.sum(axis=1)),
            "max_plan_cost_diff":_maxdiff(xcost,expected_plan_cost)}
    if has_adjust:
        adjusted=np.asarray([float(r["adjusted_grid_kwh"]) for r in rows]).reshape(334,144)
        adjust_rows=list(wb.worksheets[1].iter_rows(min_row=2,max_row=335,min_col=2,max_col=147,values_only=True))
        xadjust=np.asarray([r[:144] for r in adjust_rows],dtype=float)
        xa_total=np.asarray([r[144] for r in adjust_rows],dtype=float)
        xa_cost=np.asarray([r[145] for r in adjust_rows],dtype=float)
        expected_cost=np.asarray([float(r["settlement_cost"]) for r in daily])
        result.update({"max_adjust_diff":_maxdiff(xadjust,adjusted),
                       "max_adjust_total_diff":_maxdiff(xa_total,adjusted.sum(axis=1)),
                       "max_adjust_cost_diff":_maxdiff(xa_cost,expected_cost)})
    storage_rows=list(wb.worksheets[storage_index].iter_rows(min_row=2,max_row=2005,min_col=1,max_col=6,values_only=True))
    xcharge=np.asarray([float(r[2]) for r in storage_rows]).reshape(334,6)
    xdischarge=np.asarray([float(r[3]) for r in storage_rows]).reshape(334,6)
    xs0=np.asarray([float(storage_rows[d*6][5]) for d in range(334)])
    xs1=np.asarray([float(storage_rows[d*6+1][5]) for d in range(334)])
    result.update({"storage_rows":len(storage_rows),
        "max_storage_charge_diff":_maxdiff(xcharge,charge.reshape(334,6,24).sum(axis=2)),
        "max_storage_discharge_diff":_maxdiff(xdischarge,discharge.reshape(334,6,24).sum(axis=2)),
        "max_storage_initial_diff":_maxdiff(xs0,[float(r["soc_initial"]) for r in daily]),
        "max_storage_terminal_diff":_maxdiff(xs1,[float(r["soc_terminal"]) for r in daily])})
    erows=list(wb.worksheets[emergency_index].iter_rows(min_row=2,min_col=1,max_col=3,values_only=True))
    x_emergency=sum(float(r[2] or 0) for r in erows)
    result.update({"emergency_rows":len(erows),"emergency_sheet_total":x_emergency,
                   "emergency_csv_total":float(emergency.sum()),
                   "emergency_total_diff":abs(x_emergency-float(emergency.sum()))})
    numerical=[v for k,v in result.items() if k.startswith("max_") or k=="emergency_total_diff"]
    result["passed"]=bool(result["sheet_names_preserved"] and len(storage_rows)==2004 and max(numerical)<1e-7)
    return result
def validate_all()->dict:
    cases={
        "q1":validate_q1(),
        "q2":validate_annual(ROOT/"outputs"/"q2","result2.xlsx","q2"),
        "q3_main":validate_annual(ROOT/"outputs"/"q3","result3.xlsx","q3"),
        "q3_raw_step":validate_annual(ROOT/"outputs"/"q3"/"variants"/"raw_step","result3.xlsx","q3"),
        "q3_calibrated_step":validate_annual(ROOT/"outputs"/"q3"/"variants"/"calibrated_step","result3.xlsx","q3"),
        "q4_2":validate_annual(ROOT/"outputs"/"q4-2","result4-2.xlsx","q2"),
        "q4_3":validate_annual(ROOT/"outputs"/"q4-3","result4-3.xlsx","q3"),
    }
    result={"cases":cases,"passed":all(v["passed"] for v in cases.values())}
    (ROOT/"outputs"/"verification_summary.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    return result
if __name__=="__main__": print(json.dumps(validate_all(),ensure_ascii=False,indent=2))
