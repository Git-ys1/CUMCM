"""测试辅助：在临时目录构造被人为篡改的附件数据，用于因果性 mutation test。"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

from ..solve.io_data import DEFAULT_DATA_ROOT

ATTACHMENTS = ("附件1.xlsx", "附件2.xlsx", "附件3.xlsx", "附件4.xlsx")


def make_root(tmp: Path, source: Path = DEFAULT_DATA_ROOT) -> Path:
    """复制附件 1~4 到临时目录，返回可直接传给 run_q2/run_q3 的 data_root。"""
    tmp.mkdir(parents=True, exist_ok=True)
    for name in ATTACHMENTS:
        target = tmp / name
        if not target.exists():
            shutil.copy2(source / name, target)
    return tmp


def _mutate_sheet(path: Path, sheet: str, day_index: int, from_slot: int, factor: float) -> None:
    """把第 day_index 天（0 基）从 from_slot 开始的真实值乘以 factor。"""
    workbook = openpyxl.load_workbook(path)
    worksheet = workbook[sheet] if sheet in workbook.sheetnames else workbook.worksheets[0]
    row = day_index + 2  # 第 1 行为表头
    for col in range(from_slot + 2, worksheet.max_column + 1):
        cell = worksheet.cell(row, col)
        if cell.value is None:
            continue
        cell.value = float(cell.value) * factor
    workbook.save(path)
    workbook.close()


def mutate_attachment2_load(root: Path, day_index: int, from_slot: int = 0, factor: float = 10.0) -> None:
    _mutate_sheet(root / "附件2.xlsx", "小区负载", day_index, from_slot, factor)


def mutate_attachment2_pv(root: Path, day_index: int, from_slot: int = 0, factor: float = 10.0) -> None:
    _mutate_sheet(root / "附件2.xlsx", "光伏发电实际功率", day_index, from_slot, factor)


def mutate_attachment4_price(root: Path, day_index: int, from_slot: int = 0, factor: float = 10.0) -> None:
    _mutate_sheet(root / "附件4.xlsx", "Sheet1", day_index, from_slot, factor)


def max_abs_diff(left, right) -> float:
    import numpy as np

    return float(np.max(np.abs(np.asarray(left, dtype=float) - np.asarray(right, dtype=float))))
