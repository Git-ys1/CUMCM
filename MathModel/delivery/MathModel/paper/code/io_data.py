from __future__ import annotations
from dataclasses import dataclass
from datetime import date,datetime
from pathlib import Path
import numpy as np
import openpyxl
REPO_ROOT=Path(__file__).resolve().parents[2]
DEFAULT_DATA_ROOT=REPO_ROOT/"CUMCM2026Problems"/"C题"/"附件"
DT_HOURS=1.0/6.0; SLOTS_PER_DAY=144
@dataclass(frozen=True)
class DayData:
    labels:list[str]; price:np.ndarray; load_kw:np.ndarray; pv_kw:np.ndarray
@dataclass(frozen=True)
class AnnualData:
    dates:list[date]; load_kw:np.ndarray; pv_kw:np.ndarray
@dataclass(frozen=True)
class RollingForecastData:
    dates:list[date]; hourly_kw:np.ndarray
@dataclass(frozen=True)
class AnnualPriceData:
    dates:list[date]; price:np.ndarray
def _as_float_array(values:list[object],name:str)->np.ndarray:
    array=np.asarray(values,dtype=float)
    if array.shape!=(144,) or not np.isfinite(array).all(): raise ValueError(f"{name} invalid: {array.shape}")
    return array
def _coerce_date(value:object)->date:
    if isinstance(value,datetime): return value.date()
    if isinstance(value,date): return value
    text=str(value).strip()
    for fmt in ("%Y-%m-%d","%Y/%m/%d","%m/%d/%Y"):
        try: return datetime.strptime(text,fmt).date()
        except ValueError: pass
    raise ValueError(f"无法解析日期: {value!r}")
def read_attachment1(data_root:Path=DEFAULT_DATA_ROOT)->DayData:
    wb=openpyxl.load_workbook(Path(data_root)/"附件1.xlsx",data_only=True,read_only=True)
    rows=list(wb[wb.sheetnames[0]].iter_rows(min_row=2,max_row=145,min_col=1,max_col=4,values_only=True))
    labels=[str(r[0]) for r in rows]; price=_as_float_array([r[1] for r in rows],"price")
    load=_as_float_array([r[2] for r in rows],"load"); pv=_as_float_array([r[3] for r in rows],"pv")
    if (price<=0).any() or (load<0).any() or (pv<0).any(): raise ValueError("attachment1 bounds")
    return DayData(labels,price,load,pv)
def read_attachment2(data_root:Path=DEFAULT_DATA_ROOT)->AnnualData:
    wb=openpyxl.load_workbook(Path(data_root)/"附件2.xlsx",data_only=True,read_only=True)
    arrays=[]; date_lists=[]
    for sheet_name in wb.sheetnames[:2]:
        rows=list(wb[sheet_name].iter_rows(min_row=2,max_row=366,min_col=1,max_col=145,values_only=True))
        date_lists.append([_coerce_date(r[0]) for r in rows]); values=np.asarray([r[1:] for r in rows],dtype=float)
        if values.shape!=(365,144) or not np.isfinite(values).all() or (values<0).any(): raise ValueError(f"{sheet_name} invalid")
        arrays.append(values)
    if date_lists[0]!=date_lists[1]: raise ValueError("attachment2 date mismatch")
    return AnnualData(date_lists[0],arrays[0],arrays[1])
def read_attachment3(data_root:Path=DEFAULT_DATA_ROOT)->RollingForecastData:
    wb=openpyxl.load_workbook(Path(data_root)/"附件3.xlsx",data_only=True,read_only=True)
    rows=list(wb[wb.sheetnames[0]].iter_rows(min_row=2,max_row=1461,min_col=1,max_col=26,values_only=True))
    if len(rows)!=1460: raise ValueError("attachment3 must have 1460 forecast rows")
    dates=[]; blocks=[]; current_date=None
    expected=("0:00","6:00","12:00","18:00")
    for day in range(365):
        chunk=rows[day*4:(day+1)*4]
        current_date=_coerce_date(chunk[0][0])
        dates.append(current_date)
        issues=tuple(str(r[1]) for r in chunk)
        if issues!=expected: raise ValueError(f"attachment3 issue times invalid on {current_date}: {issues}")
        values=np.asarray([r[2:] for r in chunk],dtype=float)
        if values.shape!=(4,24) or not np.isfinite(values).all() or (values<0).any(): raise ValueError(f"attachment3 invalid {current_date}")
        blocks.append(values)
    return RollingForecastData(dates,np.stack(blocks))
def hourly_to_ten_min(hourly_kw:np.ndarray,method:str="step")->np.ndarray:
    hourly=np.asarray(hourly_kw,dtype=float)
    if hourly.shape!=(24,): raise ValueError("hourly forecast must have 24 values")
    if method=="step": return np.repeat(hourly,6)
    if method=="linear":
        return np.interp(np.arange(1,145)/6.0,np.arange(1,25),hourly,left=hourly[0],right=hourly[-1])
    raise ValueError(f"unknown downscale method: {method}")
def read_attachment4(data_root:Path=DEFAULT_DATA_ROOT)->AnnualPriceData:
    wb=openpyxl.load_workbook(Path(data_root)/"附件4.xlsx",data_only=True,read_only=True)
    rows=list(wb[wb.sheetnames[0]].iter_rows(min_row=2,max_row=366,min_col=1,max_col=145,values_only=True))
    dates=[_coerce_date(r[0]) for r in rows]; price=np.asarray([r[1:] for r in rows],dtype=float)
    if price.shape!=(365,144) or not np.isfinite(price).all() or (price<=0).any(): raise ValueError("attachment4 invalid")
    return AnnualPriceData(dates,price)
def template_path(name:str,data_root:Path=DEFAULT_DATA_ROOT)->Path:
    path=Path(data_root)/"附件5"/name
    if not path.exists(): raise FileNotFoundError(path)
    return path
