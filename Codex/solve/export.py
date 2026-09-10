from __future__ import annotations
from copy import copy
from datetime import datetime,time
from pathlib import Path
from shutil import copy2
import numpy as np
import openpyxl
from .lp_core import DispatchResult

def export_result1(template:Path,destination:Path,result:DispatchResult)->dict:
    destination.parent.mkdir(parents=True,exist_ok=True); copy2(template,destination)
    wb=openpyxl.load_workbook(destination); plan,storage=wb.worksheets[:2]
    for row,value in enumerate(result.grid,start=2):
        plan.cell(row,2,float(value)); plan.cell(row,2).number_format="0.0000"
    for block in range(6):
        lo,hi=block*24,(block+1)*24
        storage.cell(block+2,2,float(result.charge[lo:hi].sum()))
        storage.cell(block+2,3,float(result.discharge[lo:hi].sum()))
    storage.cell(2,5,6000.0); storage.cell(3,5,float(result.soc[-1])); wb.save(destination)
    check=openpyxl.load_workbook(destination,data_only=True,read_only=True)
    values=np.asarray([check.worksheets[0].cell(i,2).value for i in range(2,146)],dtype=float)
    diff=float(np.max(np.abs(values-result.grid)))
    return {"rows":len(values),"max_plan_roundtrip_diff":diff,"passed":bool(diff<1e-9)}

def _copy_row_style(ws,source_row:int,target_row:int,max_col:int)->None:
    for col in range(1,max_col+1):
        src=ws.cell(source_row,col); dst=ws.cell(target_row,col)
        if src.has_style: dst._style=copy(src._style)
        dst.alignment=copy(src.alignment); dst.protection=copy(src.protection)
    ws.row_dimensions[target_row].height=ws.row_dimensions[source_row].height

def _segments(values:np.ndarray,tolerance:float=1e-7)->list[tuple[int,int,float]]:
    result=[]; start=None
    for idx,flag in enumerate(np.r_[np.asarray(values)>tolerance,False]):
        if flag and start is None: start=idx
        elif not flag and start is not None:
            result.append((start,idx,float(np.sum(values[start:idx])))); start=None
    return result

def _slot_time(slot:int)->str:
    minutes=slot*10
    return "24:00" if minutes==1440 else f"{minutes//60}:{minutes%60:02d}"

def export_result2(template:Path,destination:Path,days:list[dict])->dict:
    if len(days)!=334: raise ValueError(f"result2 requires 334 days, got {len(days)}")
    destination.parent.mkdir(parents=True,exist_ok=True); copy2(template,destination)
    wb=openpyxl.load_workbook(destination); plan_ws,storage_ws,emergency_ws=wb.worksheets[:3]
    for idx,day in enumerate(days):
        row=idx+2; plan=day["plan"]
        plan_ws.cell(row,1,datetime.combine(day["date"],time()))
        for slot,value in enumerate(plan.grid,start=2):
            plan_ws.cell(row,slot,float(value)); plan_ws.cell(row,slot).number_format="0.0000"
        plan_ws.cell(row,146,float(np.sum(plan.grid))); plan_ws.cell(row,147,float(day["plan_cost"]))
        plan_ws.cell(row,146).number_format=plan_ws.cell(row,147).number_format="0.0000"
    target=1+6*len(days)
    if storage_ws.max_row<target: storage_ws.insert_rows(storage_ws.max_row+1,amount=target-storage_ws.max_row)
    periods=("0:00-4:00","4:00-8:00","8:00-12:00","12:00-16:00","16:00-20:00","20:00-24:00")
    for idx,day in enumerate(days):
        actual=day["execution"]
        for block in range(6):
            row=2+idx*6+block; _copy_row_style(storage_ws,2+block,row,6)
            storage_ws.cell(row,1,datetime.combine(day["date"],time()) if block==0 else None)
            storage_ws.cell(row,2,periods[block]); lo,hi=block*24,(block+1)*24
            storage_ws.cell(row,3,float(actual.charge[lo:hi].sum()))
            storage_ws.cell(row,4,float(actual.discharge[lo:hi].sum()))
            storage_ws.cell(row,5,time(0,0) if block==0 else ("24:00" if block==1 else None))
            storage_ws.cell(row,6,float(day["soc_initial"]) if block==0 else (float(actual.soc[-1]) if block==1 else None))
            for col in (3,4,6): storage_ws.cell(row,col).number_format="0.0000"
    if storage_ws.max_row>target: storage_ws.delete_rows(target+1,storage_ws.max_row-target)
    rows=[]
    for day in days:
        segments=_segments(day["execution"].emergency)
        if not segments: rows.append((datetime.combine(day["date"],time()),"无",0.0))
        else:
            for k,(start,end,total) in enumerate(segments):
                rows.append((datetime.combine(day["date"],time()) if k==0 else None,
                             f"{_slot_time(start)}-{_slot_time(end)}",total))
    target_e=1+len(rows)
    if emergency_ws.max_row<target_e: emergency_ws.insert_rows(emergency_ws.max_row+1,amount=target_e-emergency_ws.max_row)
    for row,(date_value,period,total) in enumerate(rows,start=2):
        _copy_row_style(emergency_ws,2,row,3)
        emergency_ws.cell(row,1,date_value); emergency_ws.cell(row,2,period)
        emergency_ws.cell(row,3,float(total)); emergency_ws.cell(row,3).number_format="0.0000"
    if emergency_ws.max_row>target_e: emergency_ws.delete_rows(target_e+1,emergency_ws.max_row-target_e)
    wb.save(destination)
    check=openpyxl.load_workbook(destination,data_only=True,read_only=True)
    values=np.asarray([[check.worksheets[0].cell(r,c).value for c in range(2,146)] for r in range(2,336)],dtype=float)
    expected=np.stack([d["plan"].grid for d in days]); diff=float(np.max(np.abs(values-expected)))
    last_soc=float(check.worksheets[1].cell(3+6*(len(days)-1),6).value)
    return {"delivery_days":334,"plan_rows":values.shape[0],"plan_slots":values.shape[1],
            "storage_rows":check.worksheets[1].max_row-1,"emergency_rows":check.worksheets[2].max_row-1,
            "max_plan_roundtrip_diff":diff,"last_terminal_soc":last_soc,
            "passed":bool(values.shape==(334,144) and diff<1e-9 and check.worksheets[1].max_row-1==2004)}
