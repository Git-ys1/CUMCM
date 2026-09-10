from __future__ import annotations
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import numpy as np
import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATA_ROOT = REPO_ROOT / "CUMCM2026Problems" / "C题" / "附件"
DT_HOURS = 1.0/6.0
SLOTS_PER_DAY = 144

@dataclass(frozen=True)
class DayData:
    labels:list[str]; price:np.ndarray; load_kw:np.ndarray; pv_kw:np.ndarray

@dataclass(frozen=True)
class AnnualData:
    dates:list[date]; load_kw:np.ndarray; pv_kw:np.ndarray

def _as_float_array(values:list[object],name:str)->np.ndarray:
    try: array=np.asarray(values,dtype=float)
    except (TypeError,ValueError) as exc: raise ValueError(f"{name} 含非数值单元格") from exc
    if array.shape!=(SLOTS_PER_DAY,): raise ValueError(f"{name} 应有144个时段，实际{array.shape}")
    if not np.isfinite(array).all(): raise ValueError(f"{name} 含 NaN/Inf")
    return array

def read_attachment1(data_root:Path=DEFAULT_DATA_ROOT)->DayData:
    wb=openpyxl.load_workbook(Path(data_root)/"附件1.xlsx",data_only=True,read_only=True)
    rows=list(wb[wb.sheetnames[0]].iter_rows(min_row=2,max_row=145,min_col=1,max_col=4,values_only=True))
    labels=[str(r[0]) for r in rows]
    price=_as_float_array([r[1] for r in rows],"电价")
    load=_as_float_array([r[2] for r in rows],"负载")
    pv=_as_float_array([r[3] for r in rows],"光伏")
    if (price<=0).any() or (load<0).any() or (pv<0).any(): raise ValueError("附件1存在越界值")
    return DayData(labels,price,load,pv)

def _coerce_date(value:object)->date:
    if isinstance(value,datetime): return value.date()
    if isinstance(value,date): return value
    text=str(value).strip()
    for fmt in ("%Y-%m-%d","%Y/%m/%d","%m/%d/%Y"):
        try: return datetime.strptime(text,fmt).date()
        except ValueError: pass
    raise ValueError(f"无法解析日期: {value!r}")

def read_attachment2(data_root:Path=DEFAULT_DATA_ROOT)->AnnualData:
    wb=openpyxl.load_workbook(Path(data_root)/"附件2.xlsx",data_only=True,read_only=True)
    if len(wb.sheetnames)<2: raise ValueError("附件2至少应有两个工作表")
    arrays=[]; date_lists=[]
    for sheet_name in wb.sheetnames[:2]:
        rows=list(wb[sheet_name].iter_rows(min_row=2,max_row=366,min_col=1,max_col=145,values_only=True))
        if len(rows)!=365: raise ValueError(f"{sheet_name} 日期数不是365")
        date_lists.append([_coerce_date(r[0]) for r in rows])
        values=np.asarray([r[1:] for r in rows],dtype=float)
        if values.shape!=(365,144) or not np.isfinite(values).all() or (values<0).any():
            raise ValueError(f"{sheet_name} 数据无效: {values.shape}")
        arrays.append(values)
    if date_lists[0]!=date_lists[1]: raise ValueError("负载与光伏日期不一致")
    return AnnualData(date_lists[0],arrays[0],arrays[1])

def template_path(name:str,data_root:Path=DEFAULT_DATA_ROOT)->Path:
    path=Path(data_root)/"附件5"/name
    if not path.exists(): raise FileNotFoundError(path)
    return path
